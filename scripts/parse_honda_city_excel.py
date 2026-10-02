import json
import os
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CITY_DIR = os.path.join(BASE_DIR, "public", "vehicles", "cars", "honda-cars", "honda-city", "City")

def parse_engine_workbook(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    engine_specs = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        
        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        spec_items = []
        for row in rows[1:]:
            if not row or all(cell is None for cell in row):
                continue
            param = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            e_hev = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "—"
            i_vtec = str(row[2]).strip() if len(row) > 2 and row[2] is not None else "—"
            if param:
                spec_items.append({
                    "parameter": param,
                    "eHev": e_hev,
                    "iVtec": i_vtec
                })
        
        engine_specs.append({
            "category": sheet_name,
            "headers": headers,
            "specs": spec_items
        })
    
    return engine_specs

def parse_features_workbook(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    feature_matrix = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        
        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        trim_columns = headers[1:] # ["SV", "V", "VX", "ZX", "ZX+"]
        feature_items = []
        
        for row in rows[1:]:
            if not row or all(cell is None for cell in row):
                continue
            feature_name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            if not feature_name:
                continue
            
            trims = {}
            for idx, trim_col in enumerate(trim_columns):
                val_idx = idx + 1
                val = str(row[val_idx]).strip() if val_idx < len(row) and row[val_idx] is not None else "—"
                trims[trim_col] = val
            
            feature_items.append({
                "feature": feature_name,
                "trims": trims
            })
        
        feature_matrix.append({
            "category": sheet_name,
            "trims": trim_columns,
            "features": feature_items
        })
    
    return feature_matrix

def main():
    engine_file = os.path.join(CITY_DIR, "Engine.xlsx")
    features_file = os.path.join(CITY_DIR, "Features.xlsx")
    
    print(f"Parsing {engine_file}...")
    engine_specs = parse_engine_workbook(engine_file)
    print(f"Parsed {len(engine_specs)} engine sheets.")
    for cat in engine_specs:
        print(f"  • {cat['category']}: {len(cat['specs'])} parameters")
        
    print(f"\nParsing {features_file}...")
    feature_matrix = parse_features_workbook(features_file)
    print(f"Parsed {len(feature_matrix)} feature categories.")
    for cat in feature_matrix:
        print(f"  • {cat['category']}: {len(cat['features'])} features")
        
    output_path = os.path.join(BASE_DIR, "scripts", "honda_city_parsed.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump({
            "engineSpecs": engine_specs,
            "featureMatrix": feature_matrix
        }, f, indent=2, ensure_ascii=False)
        
    print(f"\nWrote parsed data to {output_path}")

if __name__ == "__main__":
    main()
