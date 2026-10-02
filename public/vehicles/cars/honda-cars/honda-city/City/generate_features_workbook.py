import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Create workbook
wb = openpyxl.Workbook()
# Remove default initial sheet once populated
default_sheet = wb.active

# Define styling rules
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
cell_font = Font(name="Segoe UI", size=10)
bold_font = Font(name="Segoe UI", size=10, bold=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center")

thin_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

sheets_data = {
    "Exterior Design": [
        ("Blade-Eye Signature LED Lighting System", "—", "—", "—", "Y", "Y"),
        ("Signature Bi-LED Projector Headlamps with Integrated Split LED DRL & Turn Signals", "—", "—", "—", "Y", "Y"),
        ("Front Grille Connected Centre Light Bar (Position Lamp)", "—", "—", "Y", "Y", "Y"),
        ("Signature Z-Edge Wrap-around LED Tail Lamps", "Y", "Y", "Y (Clear Lens)", "Y (Clear Lens)", "Y (Clear Lens)"),
        ("LED Side Marker Lights in Tail Lamp", "—", "—", "—", "Y", "Y"),
        ("Signature Matrix-Mesh Black Painted Front Upper Grille", "Y", "Y", "Y", "Y", "Y"),
        ("Aero-Blade Front Bumper Air Curtains", "—", "—", "—", "Y (Black Painted)", "Y (Black Painted)"),
        ("Signature Matrix Mesh Front & Rear Lower Bumper Garnish", "—", "—", "—", "Y (Black Painted)", "Y (Black Painted)"),
        ("Body Coloured Trunk Lip Spoiler", "—", "—", "—", "Y", "Y"),
        ("Gloss-Black Trunk Centre Moulding", "—", "—", "—", "Y", "Y"),
        ("One-Touch Electric Sunroof with Slide/Tilt Function and Pinch Guard", "—", "—", "Y", "Y", "Y"),
        ("Steel Wheels R15 with Silver Trim Cover", "Y", "—", "—", "—", "—"),
        ("Multi-Spoke Gray Painted Alloy Wheels R15", "—", "Y", "—", "—", "—"),
        ("Aero-Blade Diamond Cut Alloy Wheels R16", "—", "—", "Y (Shark Gray & Bright-Clear Cut)", "Y (Berlina Black & Dark-Clear Cut)", "Y (Berlina Black & Dark-Clear Cut)"),
        ("Body Coloured Door Handles", "Y", "Y", "Y", "Y", "Y"),
        ("Body Coloured Shark Fin Antenna", "Y", "Y", "Y", "Y", "Y"),
        ("Body Coloured Door Mirrors", "Y", "Y", "Y", "Y", "Y"),
        ("Integrated LED Side Turn Indicator on Door Mirrors", "Y", "Y", "Y", "Y", "Y"),
        ("e:HEV Signature Rear Emblem", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Black Sash Tape on B-Pillar", "Y", "Y", "Y", "Y", "Y"),
    ],
    "Interior & Finish": [  # Excel limits sheet names to <= 31 characters
        ("Leather Upholstery with Luxurious Ivory & Black Two-Tone Interiors", "—", "—", "—", "Y", "Y"),
        ("Embossed-Fabric Upholstery with Premium Beige & Black Two-Tone Interiors", "Y", "Y", "Y", "—", "—"),
        ("Exclusive Seat Design", "—", "—", "—", "Y (Contemporary)", "Y (Sporty)"),
        ("Instrument Panel Assistant-side Garnish", "—", "Y (Piano Black)", "Y (Piano Black)", "Y (Dark Iron 3D Pattern)", "Y (Dark Iron 3D Pattern)"),
        ("Leather Shift Lever Boot with Stitch", "—", "—", "Y", "Y", "Y"),
        ("Smooth Leather Steering Wheel with Stitch", "—", "—", "—", "Y", "Y"),
        ("Hand-wrapped Leather Soft Pads (Dashboard Mid Pad, Centre Console Knee Pad, Front Centre Armrest, Door Lining - Armrest & Centre Pads)", "—", "—", "—", "Y", "Y"),
        ("Fabric Soft Pads (Front Centre Armrest, Door Lining - Armrest & Centre Pads)", "Y", "Y", "Y", "—", "—"),
        ("AC Vents Surround Finish", "—", "—", "—", "Y (Piano Black)", "Y (Piano Black)"),
        ("AC Vent Knobs Chrome Finish", "—", "Y", "Y", "Y", "Y"),
        ("Piano Black Garnish on Steering Wheel", "—", "—", "Y", "Y", "Y"),
        ("Inside Door Handle Chrome Finish", "—", "Y", "Y", "Y", "Y"),
        ("Chrome Decoration Ring for Map Lamp & Rear Reading Lamp", "—", "—", "—", "Y", "Y"),
        ("Trunk Lid Inside Lining Cover", "Y", "Y", "Y", "Y", "Y"),
        ("Ambient Illumination Light (Dashboard & Display Audio)", "—", "—", "—", "Y", "Y"),
        ("Front Footwell Ambient Illumination Light", "—", "—", "—", "Y", "Y"),
        ("Centre Console Pocket Ambient Illumination Light", "—", "—", "—", "Y", "Y"),
    ],
    "Safety and Security": [
        ("Honda SENSING (ADAS*)", "—", "Y", "Y", "Y", "Y"),
        ("Collision Mitigation Braking System (CMBS)", "—", "Y", "Y", "Y", "Y"),
        ("Adaptive Cruise Control", "—", "Y", "Y", "Y", "Y"),
        ("Low Speed Follow (LSF)", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Auto High-Beam", "—", "Y", "Y", "Y", "Y"),
        ("Lane Keeping Assist System (LKAS)", "—", "Y", "Y", "Y", "Y"),
        ("Road Departure Mitigation System", "—", "Y", "Y", "Y", "Y"),
        ("Lead Car Departure Notification System", "—", "Y", "Y", "Y", "Y"),
        ("360-degree Surround-Vision Camera (Multi-View with Guidelines & MOD)", "—", "—", "—", "Y", "Y"),
        ("Multi-Angle Rear Camera with Guidelines (Normal, Wide, Top-Down)", "—", "Y", "Y", "Y", "Y"),
        ("Rear Parking Sensors", "Y", "Y", "Y", "Y", "Y"),
        ("LaneWatch™ Camera", "—", "—", "Y", "Y", "Y"),
        ("Advanced Compatibility Engineering (ACE™) Body Structure", "Y", "Y", "Y", "Y", "Y"),
        ("6 Airbags System (Dual Front i-SRS, Front Seats i-Side & Side Curtain)", "Y", "Y", "Y", "Y", "Y"),
        ("Energy Absorbing Front Pretensioner Seatbelts with Load Limiter", "Y", "Y", "Y", "Y", "Y"),
        ("All 5 Seats 3-Point Emergency Locking Retractor (ELR) Seatbelts", "Y", "Y", "Y", "Y", "Y"),
        ("All 5 Seats Head Restraints", "Y", "Y", "Y", "Y", "Y"),
        ("ISOFIX Compatible Rear Side Seats with Lower Anchorage & Top Tether", "Y", "Y", "Y", "Y", "Y"),
        ("Anti-Lock Brake System (ABS)", "Y", "Y", "Y", "Y", "Y"),
        ("Brake Assist (BA) with Electronic Brake-force Distribution (EBD)", "Y", "Y", "Y", "Y", "Y"),
        ("Vehicle Stability Assist (VSA) with Electronic Stability & Traction Control", "Y", "Y", "Y", "Y", "Y"),
        ("Agile Handling Assist (AHA)", "Y", "Y", "Y", "Y", "Y"),
        ("Hill Start Assist", "Y", "Y", "Y", "Y", "Y"),
        ("Emergency Stop Signal", "Y", "Y", "Y", "Y", "Y"),
        ("Tyre Pressure Monitoring System (Deflation Warning System)", "Y", "Y", "Y", "Y", "Y"),
        ("AVAS (Acoustic Vehicle Alerting System at low speed EV Mode)", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("All 4 Wheels Disc Brakes", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Electric Parking Brake (EPB)", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Automatic Brake Hold System", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Automatic Headlight Control with Light Sensor", "Y", "Y", "Y", "Y", "Y"),
        ("Inside Rear View Mirror (Day/Night)", "Y (Manual)", "Y (Manual)", "Y (Manual)", "Y (Auto-Dimming)", "Y (Auto-Dimming)"),
        ("Advanced Wipers & Washer with Variable Intermittent Speed", "Y", "Y", "Y", "—", "—"),
        ("Advanced Auto-Wiper System with Rain Sensor", "—", "—", "—", "Y", "Y"),
        ("Rear Windshield Demister", "Y", "Y", "Y", "Y", "Y"),
        ("All Seats Seatbelt Indicator & Reminder", "Y", "Y", "Y", "Y", "Y"),
        ("Speed Alarm Indicator & Reminder", "Y", "Y", "Y", "Y", "Y"),
        ("Door Ajar & Trunk Open Indicator & Reminder", "Y", "Y", "Y", "Y", "Y"),
        ("Vehicle Immobilizer with Security System Alarm", "Y", "Y", "Y", "Y", "Y"),
        ("Dual Horn", "Y", "Y", "Y", "Y", "Y"),
        ("Battery Sensor", "Y", "Y", "Y", "Y", "Y"),
        ("Tyre Pressure Monitoring System (Linked via Honda Connect App)", "(Optional)#", "(Optional)#", "(Optional)#", "(Optional)#", "(Optional)#"),
        ("Front & Rear DVR with Parking Monitoring (Linked via Honda Connect App)", "(Optional)#", "(Optional)#", "(Optional)#", "(Optional)#", "(Optional)#"),
    ],
    "Comfort and Convenience": [
        ("Active Ventilation System (Front Seats)", "—", "—", "—", "—", "Y"),
        ("Honda Smart Key System with Keyless Remote (x2)", "Y", "Y", "Y", "Y", "Y"),
        ("One-Push Start/Stop Button", "Y", "Y", "Y", "Y", "Y"),
        ("Touch-Sensor Based Smart Keyless Access", "—", "Y", "Y", "Y", "Y"),
        ("Electrical Trunk Lock with Keyless Release", "Y", "Y", "Y", "Y", "Y"),
        ("Remote Engine Start", "—", "Y (CVT only)", "Y (CVT only)", "Y (CVT only)", "Y (CVT & e:HEV only)"),
        ("Walk Away Auto Lock (Customizable)", "—", "Y", "Y", "Y", "Y"),
        ("Power Windows & Sunroof Keyless Remote Open/Close", "—", "—", "—", "Y", "Y"),
        ("Follow-Me-Home / Lead-Me-To-Car Headlights (Auto Off Timer)", "Y", "Y", "Y", "Y", "Y"),
        ("Power Windows Auto-Open/Close Function with Pinch Guard", "Y (Driver only)", "Y (Driver only)", "Y (Driver only)", "Y (Driver only)", "Y (All 4 Doors)"),
        ("Auto-Folding Remote Retractable Door Mirrors (Welcome Function)", "—", "—", "—", "Y", "Y"),
        ("Power Adjustable & Folding Door Mirrors", "Y", "Y", "Y", "Y", "Y"),
        ("Auto Engine Stop/Start", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Electric AC Compressor & Coolant Pump with Heater", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Fully Automatic Climate Control with MAX COOL & Click-Feel Dials", "Y", "Y", "Y", "Y", "Y"),
        ("Rear AC Vents", "Y", "Y", "Y", "Y", "Y"),
        ("Rear Sunshade", "—", "—", "Y", "Y", "Y"),
        ("PM2.5 Dust & Pollen Cabin Filter", "Y", "Y", "Y", "Y", "Y"),
        ("Power Central Door Lock w. Driver Master Switch", "Y", "Y", "Y", "Y", "Y"),
        ("Automatic Door Locking & Unlock", "Y", "Y", "Y", "Y", "Y"),
        ("Telescopic & Tilt Steering Adjustment for Reach & Height", "Y", "Y", "Y", "Y", "Y"),
        ("Driver Seat Height Adjuster", "Y", "Y", "Y", "Y", "Y"),
        ("Adaptive Cruise Control & LKAS Switches on Steering Wheel", "—", "Y", "Y", "Y", "Y"),
        ("7-Speed Paddle Shifters (for Sequential Shifts)", "—", "Y (CVT only)", "Y (CVT only)", "Y (CVT only)", "Y (CVT only)"),
        ("Deceleration Selector Paddle Shifters (for Battery Regeneration)", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("LED Shift Lever Position Indicator", "—", "Y (CVT only)", "Y (CVT only)", "Y (CVT only)", "Y (CVT & e:HEV only)"),
        ("Easy Shift Lock Release Slot", "—", "Y (CVT only)", "Y (CVT only)", "Y (CVT only)", "Y (CVT & e:HEV only)"),
        ("One Touch Turn Signal for Lane Change Signalling", "Y", "Y", "Y", "Y", "Y"),
        ("Front Accessory Charging Ports with Lid", "Y", "Y", "Y", "Y", "Y"),
        ("Front Centre Armrest & Storage Box", "Y", "Y", "Y", "Y", "Y"),
        ("Front Console Lower Pocket for Smartphones", "Y", "Y", "Y", "Y", "Y"),
        ("Floor Console Cupholders & Utility Space for Smartphones", "Y", "Y", "Y", "Y", "Y"),
        ("Wireless Charger (Plug & Play Type)", "—", "—", "Y", "Y (MT & CVT only)", "Y (MT & CVT only)"),
        ("Wireless Charger (Centre Console Smartphone Tray)", "—", "—", "—", "—", "Y (e:HEV only)"),
        ("Rear Centre Foldable Armrest with Cupholder", "Y", "Y", "Y", "Y", "Y"),
        ("1L Bottle Holders in All Doors", "Y", "Y", "Y", "Y", "Y"),
        ("Seat Back Pockets", "—", "—", "Y (Passenger side)", "Y (Driver & Passenger)", "Y (Driver & Passenger)"),
        ("Sun-visor Vanity Mirrors", "Y (Passenger side)", "Y (Passenger side)", "Y (Driver & Passenger)", "Y (Driver & Passenger)", "Y (Driver & Passenger)"),
        ("Driver Side Coin Pocket with Lid", "Y", "Y", "Y", "Y", "Y"),
        ("Foldable Grab Handles (Soft Closing Motion)", "Y", "Y", "Y", "Y", "Y"),
        ("Front Map Lamps", "Y (Bulb)", "Y (Bulb)", "Y (Bulb)", "Y (LED)", "Y (LED)"),
        ("Rear Reading Lamps", "—", "—", "Y (Bulb)", "Y (LED)", "Y (LED)"),
        ("Interior Centre Roof Light", "Y (Bulb)", "Y (Bulb)", "Y (Bulb)", "—", "—"),
        ("Trunk Light for Cargo Area Illumination", "Y", "Y", "Y", "Y", "Y"),
    ]
}

headers = ["Feature", "SV", "V", "VX", "ZX", "ZX+"]

for sheet_title, rows in sheets_data.items():
    ws = wb.create_sheet(title=sheet_title)
    ws.append(headers)

    # Style header row
    ws.row_dimensions[1].height = 26
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_left if col_num == 1 else align_center
        cell.border = cell_border

    # Append and style data rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(list(row_data))
        ws.row_dimensions[row_idx].height = 22
        current_fill = zebra_fill if row_idx % 2 == 0 else white_fill

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = bold_font if col_idx == 1 else cell_font
            cell.fill = current_fill
            cell.border = cell_border
            cell.alignment = align_left if col_idx == 1 else align_center

    # Set column widths
    ws.column_dimensions['A'].width = 58
    for col_letter in ['B', 'C', 'D', 'E', 'F']:
        ws.column_dimensions[col_letter].width = 24

    # Freeze header row
    ws.freeze_panes = "A2"

# Remove initial default sheet
if "Sheet" in wb.sheetnames:
    del wb["Sheet"]

# Save workbook
output_filename = "Features.xlsx"
wb.save(output_filename)
print(f"Workbook '{output_filename}' generated with {len(wb.sheetnames)} sheets: {', '.join(wb.sheetnames)}")