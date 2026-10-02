import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Initialize Workbook
wb = openpyxl.Workbook()

# Define Professional Typography & Palette
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
param_font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
cell_font = Font(name="Segoe UI", size=10, color="0F172A")

align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(
    left=thin_border_side,
    right=thin_border_side,
    top=thin_border_side,
    bottom=thin_border_side
)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

headers = ["Parameter", "e:HEV (Self-Charging Strong Hybrid)", "i-VTEC (Petrol)"]

# Structured Tables Data
engine_data = {
    "Engine Specifications": [
        ("Type", "Water-Cooled Inline 4 Cylinder Atkinson cycle DOHC with VTC", "Water-Cooled Inline 4 Cylinder DOHC with VTC"),
        ("Displacement (cm³)", "1 498", "1 498"),
        ("No. of Valves", "16", "16"),
        ("Max Power (kW [PS] @ rpm)", "74 (100) @ 5 600 - 6 400", "89 [121] @ 6 600"),
        ("Max Torque (Nm @ rpm)", "131 @ 4 500 - 5 000", "145 @ 4 300"),
        ("Fuel Efficiency* (km/l)", "27.26", "17.77 (MT), 17.97 (CVT)"),
        ("Fuel Type", "Compliant with E20 (20% Ethanol blended) Petrol", "Compliant with E20 (20% Ethanol blended) Petrol"),
    ],
    "Hybrid EV System": [  # Excel limit: <= 31 characters
        ("System Combined Max Power (kW [ps])", "93 [126]", "—"),
        ("Battery Charging", "Self-Charging", "—"),
        ("Hybrid Vehicle Type", "Strong Hybrid Electric Vehicle (SHEV)", "—"),
        ("Electric Drive (Traction) Motor Type", "AC Synchronous (Permanent Magnet)", "—"),
        ("Max Power (kW [ps] @ rpm)", "80 (109) @ 3 500", "—"),
        ("Max Torque (Nm @ rpm)", "253 @ 0 - 3 000", "—"),
        ("Rechargeable Energy Storage System (ReESS) Type", "Advanced Lithium Ion High-Voltage Battery", "—"),
        ("Nominal Voltage (V)", "172.8", "—"),
    ],
    "Transmission": [
        ("Manual", "—", "6 Forward & 1 Reverse"),
        ("Automatic", "e-CVT (Electrically-coupled Continuously Variable Transmission) with Engine-Linked Clutch (Wet-type Multi-Plate)", "CVT (Continuously Variable Transmission)"),
        (
            "Automatic Drive Modes",
            "• [D (Drive) / B (Brake), ECON On/Off - Selectable Drive Modes]\n• [EV Drive Mode, Hybrid Drive & Engine Drive Mode - Auto Selection]",
            "• [D (Drive) / S (Sports) / SM (Sequential Shift Sports), ECON On/Off - Selectable Drive Modes]"
        ),
    ],
    "Dimensions and Weight": [
        ("Overall Length (mm)", "4 594", "4 594"),
        ("Overall Width (mm)", "1 748", "1 748"),
        ("Overall Height (mm)", "1 489", "1 489"),
        ("Wheelbase (mm)", "2 600", "2 600"),
        ("Front Track (mm)", "1 496", "1 496"),
        ("Rear Track (mm)", "1 485", "1 484"),
        ("Kerb Weight (kg)", "1 294", "1 123 - 1 142 (MT), 1 144 - 1 168 (CVT)"),
        ("Max. GVW (kg)", "1 700", "1 580"),
        ("Fuel Tank Capacity (Litres)", "40", "40"),
        ("Seating Capacity (Persons)", "5", "5"),
        ("Steering System", "Collapsible Electric Power Assisted", "Collapsible Electric Power Assisted"),
        ("Turning Circle Radius (m)", "5.3", "5.3"),
    ],
    "Suspension": [
        ("Front Suspension", "MacPherson Strut with Coil Spring", "MacPherson Strut with Coil Spring"),
        ("Rear Suspension", "Torsion Beam with Coil Spring", "Torsion Beam with Coil Spring"),
        ("Shock Absorbers", "Telescopic Hydraulic Nitrogen Gas-Filled", "Telescopic Hydraulic Nitrogen Gas-Filled"),
        ("Stabilizer Bar", "Torsion Bar Type (Front)", "Torsion Bar Type (Front)"),
    ],
    "Brake System": [
        ("Front Brakes", "Ventilated Disc", "Ventilated Disc"),
        ("Rear Brakes", "Disc (Solid)", "Drum"),
        ("Parking Brake Type", "Electric Parking Brake (EPB) with Automatic Brake Hold", "Mechanical Handbrake (Lever Type)"),
        ("Assistance Systems", "ABS with EBD, Brake Assist (BA), Regenerative Braking System", "ABS with EBD, Brake Assist (BA)"),
    ],
    "Tyres and Wheels": [
        ("Wheel Type", "16x6J Alloy", "15x6J Steel & Alloy, 16x6 J Alloy"),
        ("Tyre Type & Size", "Tubeless Radial 185/55 R16 87H", "Tubeless Radial 185/60 R15 84H, 185/55 R16 83H"),
        ("Spare Wheel Type", "15x4 T Steel, T135/80D15 100M (Compact Spare Tyre)", "Steel Wheel, 185/60R15 84H"),
    ]
}

# Populate Workbook
for sheet_title, rows in engine_data.items():
    ws = wb.create_sheet(title=sheet_title)
    ws.append(headers)

    # Style Header Row
    ws.row_dimensions[1].height = 28
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_left if col_num == 1 else align_center
        cell.border = cell_border

    # Append & Style Data Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(list(row_data))
        
        # Adjust row height based on multiline text
        has_multiline = any("\n" in str(val) for val in row_data)
        ws.row_dimensions[row_idx].height = 42 if has_multiline else 22
        
        current_fill = zebra_fill if row_idx % 2 == 0 else white_fill

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = param_font if col_idx == 1 else cell_font
            cell.fill = current_fill
            cell.border = cell_border
            cell.alignment = align_left if col_idx == 1 else align_center

    # Column Widths
    ws.column_dimensions['A'].width = 46
    ws.column_dimensions['B'].width = 48
    ws.column_dimensions['C'].width = 48

    # Freeze Header Pane
    ws.freeze_panes = "A2"

# Remove initial blank sheet
if "Sheet" in wb.sheetnames:
    del wb["Sheet"]

# Save Output
output_filename = "Engine.xlsx"
wb.save(output_filename)
print(f"Workbook '{output_filename}' generated successfully with {len(wb.sheetnames)} worksheets:")
for name in wb.sheetnames:
    print(f"  • {name}")