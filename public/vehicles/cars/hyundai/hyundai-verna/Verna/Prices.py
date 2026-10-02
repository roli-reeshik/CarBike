import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==============================================================================
# TARGET DIRECTORY SETUP
# ==============================================================================
TARGET_DIR = r"D:\Apps\CarBikeKhardo_29_09_2026\public\vehicles\cars\hyundai\hyundai-verna\Verna"
os.makedirs(TARGET_DIR, exist_ok=True)
DEST_FILE = os.path.join(TARGET_DIR, "Price-Ex-ShowRoom.xlsx")

# ==============================================================================
# TYPOGRAPHY & DESIGN PALETTE
# ==============================================================================
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
title_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")   # Slate 800

cell_font = Font(name="Segoe UI", size=10, color="0F172A")
bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
price_font = Font(name="Segoe UI", size=10, bold=True, color="047857")  # Emerald 700
note_font = Font(name="Segoe UI", size=9, italic=True, color="475569")

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

thin_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

def style_header_row(ws, row_idx, num_cols):
    ws.row_dimensions[row_idx].height = 26
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = cell_border

def apply_table_styles(ws, start_row, end_row, num_cols):
    for r in range(start_row, end_row + 1):
        ws.row_dimensions[r].height = 22
        current_fill = zebra_fill if r % 2 == 0 else white_fill
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = current_fill
            cell.border = cell_border
            if c in (1, 2, 7):
                cell.alignment = align_left
            else:
                cell.alignment = align_center

            if c == 1:
                cell.font = bold_font
            elif c == 6:
                cell.font = price_font
                cell.alignment = align_right
                cell.number_format = '₹ #,##,##0'
            elif c == 5:
                cell.font = price_font
                cell.alignment = align_right
            else:
                cell.font = cell_font

# ==============================================================================
# DATA STRUCTURES
# ==============================================================================
MATTE_NOTE = "Note: Matte Colour Option is only available in select Mono Tone Trims at an additional cost of INR 15,000 (in Ex-Showroom Price)."

COLUMNS = ["Model", "Variant Description", "Powertrain", "Transmission", "Ex-Showroom Price (INR)", "Numeric Price", "City / State"]

MT_VARIANTS = [
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 2", "1.5 l MPi Petrol", "MT", "₹ 10 99 200", 1099200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 4", "1.5 l MPi Petrol", "MT", "₹ 12 31 200", 1231200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 6", "1.5 l MPi Petrol", "MT", "₹ 13 25 200", 1325200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 6 DT", "1.5 l MPi Petrol", "MT", "₹ 13 40 200", 1340200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 6+", "1.5 l MPi Petrol", "MT", "₹ 13 87 200", 1387200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 6+ DT", "1.5 l MPi Petrol", "MT", "₹ 14 02 200", 1402200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 8", "1.5 l MPi Petrol", "MT", "₹ 14 94 200", 1494200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol MT HX 8 DT", "1.5 l MPi Petrol", "MT", "₹ 15 09 200", 1509200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 Turbo GDi Petrol MT HX 8", "1.5 l Turbo GDi Petrol", "MT", "₹ 16 34 200", 1634200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 Turbo GDi Petrol MT HX 8 DT", "1.5 l Turbo GDi Petrol", "MT", "₹ 16 49 200", 1649200, "New Delhi, Delhi")
]

IVT_VARIANTS = [
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 6", "1.5 l MPi Petrol", "IVT", "₹ 14 46 200", 1446200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 6 DT", "1.5 l MPi Petrol", "IVT", "₹ 14 61 200", 1461200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 6+", "1.5 l MPi Petrol", "IVT", "₹ 15 03 200", 1503200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 6+ DT", "1.5 l MPi Petrol", "IVT", "₹ 15 18 200", 1518200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 8", "1.5 l MPi Petrol", "IVT", "₹ 16 15 200", 1615200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 8 DT", "1.5 l MPi Petrol", "IVT", "₹ 16 30 200", 1630200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 10", "1.5 l MPi Petrol", "IVT", "₹ 17 21 200", 1721200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 MPi Petrol IVT HX 10 DT", "1.5 l MPi Petrol", "IVT", "₹ 17 36 200", 1736200, "New Delhi, Delhi")
]

DCT_VARIANTS = [
    ("All-New Verna", "VERNA 1.5 Turbo GDi Petrol DCT HX 8", "1.5 l Turbo GDi Petrol", "DCT", "₹ 17 63 200", 1763200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 Turbo GDi Petrol DCT HX 8 DT", "1.5 l Turbo GDi Petrol", "DCT", "₹ 17 78 200", 1778200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 Turbo GDi Petrol DCT HX 10", "1.5 l Turbo GDi Petrol", "DCT", "₹ 18 31 200", 1831200, "New Delhi, Delhi"),
    ("All-New Verna", "VERNA 1.5 Turbo GDi Petrol DCT HX 10 DT", "1.5 l Turbo GDi Petrol", "DCT", "₹ 18 46 200", 1846200, "New Delhi, Delhi")
]

ALL_VARIANTS = MT_VARIANTS + IVT_VARIANTS + DCT_VARIANTS

STATES_AND_UT = [
    "Andaman & Nicobar Islands", "Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar",
    "Chandigarh", "Chhattisgarh", "Dadra Nagar Haveli", "Goa", "Gujarat",
    "Haryana", "Himachal Pradesh", "Jammu & Kashmir", "Jharkhand", "Karnataka",
    "Kerala", "Madhya Pradesh", "Maharashtra", "Manipur", "Meghalaya",
    "Mizoram", "New Delhi", "Odisha", "Punjab", "Rajasthan",
    "Sikkim", "Tamil Nadu", "Telangana", "Uttar Pradesh", "Uttarakhand", "West Bengal"
]

# ==============================================================================
# WORKBOOK GENERATION
# ==============================================================================
def create_price_workbook():
    wb = openpyxl.Workbook()

    sheets = [
        ("All Variants Summary", ALL_VARIANTS),
        ("Manual (MT)", MT_VARIANTS),
        ("Automatic (IVT)", IVT_VARIANTS),
        ("Dual Clutch (DCT)", DCT_VARIANTS)
    ]

    for title, dataset in sheets:
        ws = wb.create_sheet(title=title)
        ws.append(COLUMNS)
        style_header_row(ws, 1, len(COLUMNS))

        for row in dataset:
            ws.append(list(row))

        end_r = 1 + len(dataset)
        apply_table_styles(ws, 2, end_r, len(COLUMNS))

        # Add Note
        note_row = end_r + 2
        ws.cell(row=note_row, column=1, value=MATTE_NOTE).font = note_font
        ws.row_dimensions[note_row].height = 20

        # Width sizing
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 44
        ws.column_dimensions['C'].width = 24
        ws.column_dimensions['D'].width = 16
        ws.column_dimensions['E'].width = 25
        ws.column_dimensions['F'].width = 18
        ws.column_dimensions['G'].width = 20
        ws.freeze_panes = "A2"

    # Reference Sheet: States & Territories
    ws_states = wb.create_sheet(title="Supported States & Cities")
    ws_states.append(["S.No.", "State / Union Territory", "Primary City Configured"])
    style_header_row(ws_states, 1, 3)

    for idx, state_name in enumerate(STATES_AND_UT, start=1):
        city = "New Delhi" if state_name == "New Delhi" else f"{state_name} Dealer Network"
        ws_states.append([idx, state_name, city])

    apply_table_styles(ws_states, 2, 1 + len(STATES_AND_UT), 3)
    ws_states.column_dimensions['A'].width = 10
    ws_states.column_dimensions['B'].width = 34
    ws_states.column_dimensions['C'].width = 30
    ws_states.freeze_panes = "A2"

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(DEST_FILE)
    print(f"Created: {DEST_FILE}")

if __name__ == "__main__":
    create_price_workbook()