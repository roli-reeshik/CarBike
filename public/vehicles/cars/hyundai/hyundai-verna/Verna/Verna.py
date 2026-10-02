import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==============================================================================
# TARGET DIRECTORY SETUP
# ==============================================================================
TARGET_DIR = r"D:\Apps\CarBikeKhardo_29_09_2026\public\vehicles\cars\hyundai\hyundai-verna\Verna"
os.makedirs(TARGET_DIR, exist_ok=True)

# ==============================================================================
# COMMON STYLING DEFINITIONS
# ==============================================================================
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Slate 900
cell_font = Font(name="Segoe UI", size=10, color="0F172A")
bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
footnote_font = Font(name="Segoe UI", size=9, italic=True, color="475569")

align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

def apply_table_styles(ws, start_row, end_row, num_cols):
    for r in range(start_row, end_row + 1):
        ws.row_dimensions[r].height = 22
        current_fill = zebra_fill if r % 2 == 0 else white_fill
        for c in range(1, num_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = bold_font if c == 1 else cell_font
            cell.fill = current_fill
            cell.border = cell_border
            cell.alignment = align_left if c == 1 else align_center

# ==============================================================================
# 1. CREATE "Features.xlsx"
# ==============================================================================
def create_features_workbook():
    wb = openpyxl.Workbook()
    headers = ["Feature", "HX 2", "HX 4", "HX 6", "HX 6+", "HX 8", "HX 10"]

    features_data = {
        "Safety": [
            ("Airbag - Driver & passenger", "S", "S", "S", "S", "S", "S"),
            ("Airbag - Side & curtain", "S", "S", "S", "S", "S", "S"),
            ("Airbag - Center side", "—", "—", "—", "—", "—", "S"),
            ("Hill-start assist control (HAC)", "S", "S", "S", "S", "S", "S"),
            ("Electronic stability control (ESC)", "S", "S", "S", "S", "S", "S"),
            ("Vehicle stability management (VSM)", "S", "S", "S", "S", "S", "S"),
            ("ABS (Anti-lock braking system) with EBD", "S", "S", "S", "S", "S", "S"),
            ("Emergency stop signal (ESS)", "S", "S", "S", "S", "S", "S"),
            ("Parking assist - Rear parking sensors", "S", "S", "S", "S", "S", "S"),
            ("Parking assist - Rear view camera", "—", "—", "S", "S", "S", "S"),
            ("Parking assist - Front parking sensors", "—", "—", "S", "S", "S", "S"),
            ("Driver rear view monitor (DRVM)", "—", "—", "S", "S", "S", "S"),
            ("Tyre pressure monitoring system (TPMS) highline", "—", "S", "S", "S", "S", "S"),
            ("Inside rear view mirror - Day and night mirror", "S", "S", "—", "—", "—", "—"),
            ("Inside rear view mirror - Electro chromic mirror (ECM)", "—", "—", "S", "S", "—", "—"),
            ("Inside rear view mirror - ECM with telematic switches", "—", "—", "—", "—", "S", "S"),
            ("Electric parking brake (EPB)", "—", "—", "—", "—", "DCT only", "S"),
            ("Automatic headlamps", "S", "S", "S", "S", "S", "S"),
            ("Headlamp escort function", "S", "S", "S", "S", "S", "S"),
            ("Central locking", "S", "S", "S", "S", "S", "S"),
            ("Impact sensing auto door unlock", "S", "S", "S", "S", "S", "S"),
            ("Speed sensing auto door lock", "S", "S", "S", "S", "S", "S"),
            ("Keyless entry - Foldable key", "S", "S", "—", "—", "—", "—"),
            ("Keyless entry - Smart key with push button start", "—", "—", "S", "S", "S", "S"),
            ("Rear defogger with timer", "S", "S", "S", "S", "S", "S"),
            ("Seatbelt reminder (all seats)", "S", "S", "S", "S", "S", "S"),
            ("Height adjustable front seat belts", "—", "—", "S", "S", "S", "S"),
            ("Immobilizer", "S", "S", "S", "S", "S", "S"),
            ("Dual horn", "S", "S", "S", "S", "S", "S"),
            ("ISOFIX", "S", "S", "S", "S", "S", "S"),
            ("Rear disc brakes", "—", "—", "—", "—", "DCT Only", "S"),
            ("Burglar alarm", "S", "S", "S", "S", "S", "S"),
            ("Dashcam", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Forward Collision Warning (FCW)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Forward Collision - Avoidance Assist - Car (FCA-Car)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Forward Collision - Avoidance Assist - Pedestrian (FCA-Ped)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Forward Collision - Avoidance Assist - Cycle (FCA-Cyl)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Forward Collision - Avoidance Assist - Junction Turning (FCA-JT)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Forward Collision - Avoidance Assist - Direct Oncoming (FCA-DO)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Blind-spot Collision Warning (BCW)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Blind-spot Collision - Avoidance Assist (BCA)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Lane Keeping Assist (LKA)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Lane Departure Warning (LDW)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Driver Attention Warning (DAW)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Safe Exit Warning (SEW)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Smart Cruise Control With Stop & Go (SCC with S&G)", "—", "—", "—", "—", "DCT only", "S"),
            ("ADAS - Lane Following Assist (LFA)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - High Beam Assist (HBA)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Leading Vehicle Departure Alert (LVDA)", "—", "—", "—", "—", "DCT only", "S"),
            ("ADAS - Rear Cross - Traffic Collision Warning (RCCW)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("ADAS - Rear Cross - Traffic Collision - Avoidance Assist (RCCA)", "—", "—", "—", "—", "1.5 T-PL Only", "S"),
            ("Surround view monitor (SVM)", "—", "—", "—", "—", "—", "S"),
            ("Blind-spot view monitor (BVM)", "—", "—", "—", "—", "—", "S")
        ],
        "Exterior": [
            ("Headlamps - Projector headlamps", "S", "S", "—", "—", "—", "—"),
            ("Headlamps - Dual LED projector headlamps", "—", "—", "S", "S", "S", "S"),
            ("LED positioning lamp & DRLs", "—", "S", "S", "S", "S", "S"),
            ("LED tail lamps", "—", "S", "S", "S", "S", "S"),
            ("Rear spoiler", "—", "S", "S", "S", "S", "S"),
            ("Black chrome radiator grille", "S", "S", "S", "S", "S", "S"),
            ("Satin chrome window belt line", "—", "—", "S", "S", "S", "S"),
            ("Outside door handles - Body colored", "S", "S", "—", "—", "—", "—"),
            ("Outside door handles - Satin chrome", "—", "—", "S", "S", "S", "S"),
            ("Shark fin antenna", "—", "S", "S", "S", "S", "S"),
            ("Wheels - R15 (D=380.2 mm) Steel wheel with wheel cover", "S", "—", "—", "—", "—", "—"),
            ("Wheels - R15 (D=380.2 mm) Silver alloys", "—", "S", "—", "—", "—", "—"),
            ("Wheels - R16 (D=405.6 mm) Diamond cut alloys", "—", "—", "S", "S", "1.5 PL only", "1.5 PL only"),
            ("Wheels - R16 (D=405.6 mm) Dark grey alloys", "—", "—", "—", "—", "1.5 T-PL Only", "1.5 T-PL Only"),
            ("Red front brake calipers", "—", "—", "—", "—", "1.5 T-PL Only", "1.5 T-PL Only")
        ],
        "Interior": [
            ("Interior color theme - Dual tone", "S", "S", "S", "S", "1.5 PL only", "1.5 PL only"),
            ("Interior color theme - Black with red accents", "—", "—", "—", "—", "1.5 T-PL Only", "1.5 T-PL Only"),
            ("Seat upholstery - Cloth", "S", "S", "S", "—", "—", "—"),
            ("Seat upholstery - Leatherette (^^)", "—", "—", "—", "S", "S", "S"),
            ("D-Cut steering wheel", "S", "S", "S", "S", "S", "S"),
            ("Driver seat adjust - Manual height adjust", "S", "S", "S", "S", "—", "—"),
            ("Driver seat adjust - Electric 8-way", "—", "—", "—", "—", "S", "S"),
            ("Driver seat adjust - Memory function", "—", "—", "—", "—", "S", "S"),
            ("Passenger seat adjust - Electric 4-way", "—", "—", "—", "—", "—", "S"),
            ("Passenger seat adjust - Electric walk-in device", "—", "—", "—", "—", "—", "S"),
            ("Leatherette pack - Steering wheel", "—", "—", "S", "S", "S", "S"),
            ("Leatherette pack - Gear knob", "—", "—", "S", "S", "S", "S"),
            ("Sliding front center armrest with storage", "—", "S", "S", "S", "S", "S"),
            ("Rear center armrest with cup holders", "S", "S", "S", "S", "S", "S"),
            ("Height adjustable headrest - Front", "S", "S", "S", "S", "S", "S"),
            ("Height adjustable headrest - Rear", "—", "—", "S", "S", "S", "S")
        ],
        "Infotainment & connectivity": [
            ("Infotainment system - 20.32 cm (8.0\") Touchscreen infotainment system", "—", "S", "S", "S", "—", "—"),
            ("Infotainment system - 26.03 cm (10.25\") HD audio video navigation system", "—", "—", "—", "—", "S", "S"),
            ("Digital cluster - Digital cluster with color TFT MID", "—", "S", "S", "S", "S", "—"),
            ("Digital cluster - 26.03 cm (10.25\") Multi display digital cluster", "—", "—", "—", "—", "—", "S"),
            ("Hyundai Blue Link (Connected car technology)", "—", "—", "—", "—", "S", "S"),
            ("Wireless Apple CarPlay & Android Auto**", "—", "S", "S", "S", "#", "#"),
            ("Voice recognition**", "—", "S", "S", "S", "S", "S"),
            ("Front & rear speakers", "—", "S", "S", "S", "S", "S"),
            ("Front tweeter", "—", "—", "S", "S", "S", "S"),
            ("Bluetooth connectivity", "—", "S", "S", "S", "S", "S"),
            ("Bose premium sound 8 speaker system", "—", "—", "—", "—", "S", "S"),
            ("Steering wheel with audio & bluetooth controls", "—", "S", "S", "S", "S", "S")
        ],
        "Comfort & convenience": [
            ("Rear window sunshade", "—", "—", "S", "S", "S", "S"),
            ("Front ventilated seats", "—", "—", "—", "S", "S", "S"),
            ("Idle stop & go (ISG)", "—", "S", "S", "S", "S", "S"),
            ("Paddle shifters", "—", "—", "iVT", "iVT", "iVT & DCT", "S"),
            ("Electric sunroof", "—", "S", "S", "S", "S", "S"),
            ("Smart trunk", "—", "—", "S", "S", "S", "S"),
            ("Rain sensing wipers", "—", "—", "—", "—", "S", "S"),
            ("Drive mode select", "—", "—", "iVT", "iVT", "S (##)", "S"),
            ("Power window (Front & Rear)", "S", "S", "S", "S", "S", "S"),
            ("Air conditioning - Manual", "S", "—", "—", "—", "—", "—"),
            ("Air conditioning - Fully automatic temperature control", "—", "S", "S", "S", "S", "S"),
            ("Tilt steering", "S", "S", "S", "S", "S", "S"),
            ("Telescopic steering", "—", "S", "S", "S", "S", "S"),
            ("Rear AC vents", "—", "S", "S", "S", "S", "S"),
            ("Smartphone Wireless charger^", "—", "—", "S", "S", "S", "S"),
            ("Ambient light (dashboard & door trims)", "—", "—", "S", "S", "S", "S"),
            ("Glovebox cooling", "—", "S", "S", "S", "S", "S"),
            ("Cruise control", "—", "S", "S", "S", "S", "S"),
            ("Type-C USB Port (front & rear)", "S", "S", "S", "S", "S", "S"),
            ("Power outlet", "S", "S", "S", "S", "S", "S"),
            ("LED map Lamp", "—", "S", "S", "S", "S", "S"),
            ("LED room Lamp", "—", "S", "S", "S", "S", "S"),
            ("Metal pedals", "—", "—", "—", "—", "1.5 T-PL Only", "1.5 T-PL Only"),
            ("Luggage lamp", "S", "S", "S", "S", "S", "S"),
            ("Outside mirrors - Electrically adjustable", "S", "S", "S", "S", "S", "S"),
            ("Outside mirrors - Electric folding", "—", "—", "S", "S", "S", "S")
        ]
    }

    notes = [
        "S - Standard | — Not Available",
        "^^ Leatherette",
        "** Works with select Android & Apple smartphones only",
        "^ Works with compatible smartphones only",
        "# Through wired to wireless adapter",
        "## 1.5 PL IVT & 1.5T-PL (MT & DCT)",
        "Disclaimer: HX 8 DCT & HX 10 (iVT & DCT) will have smart cruise control S&G (SCC with S&G) instead of Cruise Control."
    ]

    for sheet_title, rows in features_data.items():
        ws = wb.create_sheet(title=sheet_title)
        ws.append(headers)

        ws.row_dimensions[1].height = 26
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_left if c == 1 else align_center
            cell.border = cell_border

        for row_idx, rdata in enumerate(rows, start=2):
            ws.append(list(rdata))

        end_row = 1 + len(rows)
        apply_table_styles(ws, 2, end_row, len(headers))

        start_note_row = end_row + 2
        for n_offset, note_text in enumerate(notes):
            n_row = start_note_row + n_offset
            ws.cell(row=n_row, column=1, value=note_text).font = footnote_font
            ws.row_dimensions[n_row].height = 18

        ws.column_dimensions['A'].width = 54
        for ch in ['B', 'C', 'D', 'E', 'F', 'G']:
            ws.column_dimensions[ch].width = 17
        ws.freeze_panes = "A2"

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    dest_file = os.path.join(TARGET_DIR, "Features.xlsx")
    wb.save(dest_file)
    print(f"Created: {dest_file}")

# ==============================================================================
# 2. CREATE "Engine.xlsx"
# ==============================================================================
def create_engine_workbook():
    wb = openpyxl.Workbook()

    # Sheet 1: Engine & Trim Plan
    ws_plan = wb.create_sheet(title="Engine & Trim Plan")
    plan_headers = ["Engine Variant", "HX 2", "HX 4", "HX 6", "HX 6+", "HX 8", "HX 10"]
    ws_plan.append(plan_headers)

    ws_plan.row_dimensions[1].height = 26
    for c in range(1, len(plan_headers) + 1):
        cell = ws_plan.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_left if c == 1 else align_center
        cell.border = cell_border

    plan_rows = [
        ("1.5 l MPi Petrol", "MT", "MT", "MT / iVT", "MT / iVT", "MT / iVT", "iVT"),
        ("1.5 l Turbo GDi Petrol", "—", "—", "—", "—", "MT / DCT", "DCT")
    ]
    for rdata in plan_rows:
        ws_plan.append(list(rdata))

    apply_table_styles(ws_plan, 2, 1 + len(plan_rows), len(plan_headers))
    ws_plan.column_dimensions['A'].width = 30
    for ch in ['B', 'C', 'D', 'E', 'F', 'G']:
        ws_plan.column_dimensions[ch].width = 16
    ws_plan.freeze_panes = "A2"

    # Sheet 2: Engine Specifications
    ws_specs = wb.create_sheet(title="Engine Specifications")
    spec_headers = ["Parameter", "1.5 l MPi Petrol", "1.5 l Turbo GDi Petrol"]
    ws_specs.append(spec_headers)

    ws_specs.row_dimensions[1].height = 26
    for c in range(1, len(spec_headers) + 1):
        cell = ws_specs.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_left if c == 1 else align_center
        cell.border = cell_border

    specs_rows = [
        ("Engine Type", "Multi-Point Injection (MPi) 4-Cylinder", "Turbocharged Gasoline Direct Injection (GDi) 4-Cylinder"),
        ("Displacement (cm³)", "1 497", "1 482"),
        ("Max Power (kW [PS] @ rpm)", "84.6 kW [115 PS] @ 6 300 rpm", "117.5 kW [160 PS] @ 5 500 rpm"),
        ("Max Torque (Nm @ rpm)", "143.8 Nm @ 4 500 rpm", "253 Nm @ 1 500 - 3 500 rpm"),
        ("Transmission Available", "6-Speed MT / Intelligent Variable Transmission (iVT)", "6-Speed MT / 7-Speed Dual Clutch Transmission (DCT)"),
        ("Fuel Type", "Petrol (E20 Compliant)", "Petrol (E20 Compliant)"),
        ("Idle Stop & Go (ISG)", "Supported (from HX 4 upwards)", "Supported (from HX 8 upwards)")
    ]
    for rdata in specs_rows:
        ws_specs.append(list(rdata))

    apply_table_styles(ws_specs, 2, 1 + len(specs_rows), len(spec_headers))
    ws_specs.column_dimensions['A'].width = 34
    ws_specs.column_dimensions['B'].width = 46
    ws_specs.column_dimensions['C'].width = 48
    ws_specs.freeze_panes = "A2"

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    dest_file = os.path.join(TARGET_DIR, "Engine.xlsx")
    wb.save(dest_file)
    print(f"Created: {dest_file}")

# ==============================================================================
# MAIN EXECUTION
# ==============================================================================
if __name__ == "__main__":
    create_features_workbook()
    create_engine_workbook()